"""
Script: convert_artikel_to_word.py
Deskripsi: Mengonversi naskah Artikel Ilmiah IndoBERT vs BiLSTM (Markdown)
           ke format Microsoft Word (.docx) berstandar jurnal ilmiah nasional/internasional
           (Times New Roman, spasi 1.15, tabel profesional, dan seluruh 7 gambar resolusi tinggi disematkan).
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
ARTIKEL_DIR = os.path.join(BASE_DIR, "07_Artikel_Ilmiah")
MD_FILE = os.path.join(ARTIKEL_DIR, "Artikel_Ilmiah_IndoBERT_vs_BiLSTM.md")
DOCX_FILE = os.path.join(ARTIKEL_DIR, "Artikel_Ilmiah_IndoBERT_vs_BiLSTM.docx")
IMAGES_DIR = os.path.join(ARTIKEL_DIR, "images")

print("[*] Membaca naskah Markdown Artikel Ilmiah...")
with open(MD_FILE, "r", encoding="utf-8") as f:
    md_content = f.read()

doc = docx.Document()

# Page Setup: Margin 1 inch (2.54 cm)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_table(table, num_cols):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    
    col_widths = None
    if num_cols == 5:
        col_widths = [Inches(1.3), Inches(2.5), Inches(1.3), Inches(1.4), Inches(1.3)]
    elif num_cols == 6:
        col_widths = [Inches(1.8), Inches(1.0), Inches(1.1), Inches(1.0), Inches(1.4), Inches(1.2)]
    elif num_cols == 7:
        col_widths = [Inches(1.5), Inches(1.2), Inches(0.7), Inches(0.8), Inches(0.8), Inches(1.1), Inches(0.9)]
    elif num_cols == 8:
        col_widths = [Inches(1.3), Inches(1.7), Inches(0.6), Inches(0.6), Inches(0.6), Inches(0.6), Inches(0.6), Inches(0.7)]

    # Header Row
    for c_idx, cell in enumerate(table.rows[0].cells):
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(r'<w:shd %s w:fill="1E3A8A"/>' % nsdecls('w')))
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        if col_widths and c_idx < len(col_widths):
            cell.width = col_widths[c_idx]
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        shd_color = "F1F5F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), shd_color)))
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
            for p in cell.paragraphs:
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                elif c_idx == 1 and num_cols == 5:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(30, 41, 59)

def format_inline_runs(paragraph, text, default_size=Pt(11), default_italic=False, default_color=None):
    # Regex untuk **bold**, *italic*, dan normal
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = "Times New Roman"
        run.font.size = default_size
        if default_color:
            run.font.color.rgb = default_color
        else:
            run.font.color.rgb = RGBColor(30, 41, 59)
            
        if token.startswith("**") and token.endswith("**") and len(token) > 4:
            run.text = token[2:-2]
            run.font.bold = True
            run.font.italic = default_italic
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run.text = token[1:-1]
            run.font.italic = True
        elif token.startswith("`") and token.endswith("`") and len(token) > 2:
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(10)
            run.font.color.rgb = RGBColor(194, 24, 91)
        else:
            run.text = token
            run.font.italic = default_italic

lines = md_content.splitlines()
i = 0
in_daftar_pustaka = False

while i < len(lines):
    line = lines[i].rstrip()
    
    if not line:
        i += 1
        continue

    # 1. Judul Artikel (# )
    if line.startswith("# "):
        title_text = line[2:].strip()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(title_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy
        i += 1
        continue

    # 2. Penulis & Afiliasi
    if line.startswith("**Nimas Ayu"):
        author_text = line.replace("**", "").strip()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(author_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11.5)
        run.font.bold = True
        i += 1
        continue

    if "Badan Pusat Statistik" in line:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(line.strip())
        run.font.name = "Times New Roman"
        run.font.size = Pt(10)
        run.font.italic = True
        i += 1
        continue

    if "Korespondensi:" in line:
        clean_text = line.replace("*", "").strip()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(16)
        run = p.add_run(clean_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(9.5)
        run.font.italic = True
        run.font.color.rgb = RGBColor(71, 85, 105)
        i += 1
        continue

    if line.strip() == "---":
        i += 1
        continue

    # 3. Heading Utama (## )
    if line.startswith("## "):
        h2_text = line[3:].strip()
        if "DAFTAR PUSTAKA" in h2_text.upper():
            in_daftar_pustaka = True

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        
        if h2_text.startswith("ABSTRAK") or h2_text.startswith("ABSTRACT"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(h2_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(h2_text)
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)
        i += 1
        continue

    # 4. Sub-heading (### )
    if line.startswith("### "):
        h3_text = line[4:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(h3_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 23, 42)
        i += 1
        continue

    # 5. Gambar (![caption](images/filename.png))
    img_match = re.match(r'!\[(.*?)\]\((.*?)\)', line.strip())
    if img_match:
        caption = img_match.group(1)
        rel_img_path = img_match.group(2)
        img_filename = os.path.basename(rel_img_path)
        full_img_path = os.path.join(IMAGES_DIR, img_filename)
        
        if os.path.exists(full_img_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(12)
            p_img.paragraph_format.space_after = Pt(4)
            run_img = p_img.add_run()
            
            # Sesuaikan lebar gambar berdasarkan tipe gambar
            if "confusion_matrix" in img_filename:
                run_img.add_picture(full_img_path, width=Inches(6.4))
            elif "alur_metodologi" in img_filename or "arsitektur" in img_filename:
                run_img.add_picture(full_img_path, width=Inches(6.2))
            else:
                run_img.add_picture(full_img_path, width=Inches(5.8))
            
            # Caption Gambar
            # Jika baris berikutnya adalah teks caption italic (misal "*Gambar 1...*"), lewati agar tidak dobel
            if i + 1 < len(lines) and lines[i+1].strip().startswith("*Gambar"):
                caption_text = lines[i+1].strip().replace("*", "")
                i += 1
            else:
                caption_text = caption
                
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(12)
            run_cap = p_cap.add_run(caption_text)
            run_cap.font.name = "Times New Roman"
            run_cap.font.size = Pt(9.5)
            run_cap.font.italic = True
            run_cap.font.color.rgb = RGBColor(71, 85, 105)
        else:
            print(f"[!] Gambar tidak ditemukan: {full_img_path}")
        i += 1
        continue

    # 6. Judul Tabel (*Tabel X...*)
    if line.strip().startswith("*Tabel") and line.strip().endswith("*"):
        tbl_title = line.strip().replace("*", "")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(tbl_title)
        run.font.name = "Times New Roman"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(15, 23, 42)
        i += 1
        continue

    # 7. Parsing Tabel Markdown
    if line.strip().startswith("|") and line.strip().endswith("|"):
        table_lines = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            t_line = lines[i].strip()
            # Abaikan separator line (| :--- | :---: |)
            if not re.match(r'^\|[\s\-:]+\|$', t_line) and not all(c in '| -:' for c in t_line):
                table_lines.append(t_line)
            i += 1

        if table_lines:
            header_cells = [c.strip() for c in table_lines[0].split('|')[1:-1]]
            num_cols = len(header_cells)
            doc_table = doc.add_table(rows=len(table_lines), cols=num_cols)
            
            for r_idx, t_line in enumerate(table_lines):
                cells_data = [c.strip() for c in t_line.split('|')[1:-1]]
                for c_idx in range(min(num_cols, len(cells_data))):
                    cell_text = cells_data[c_idx].replace("**", "")
                    doc_table.rows[r_idx].cells[c_idx].text = cell_text

            style_table(doc_table, num_cols)
            # Spasi setelah tabel
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_before = Pt(0)
            p_space.paragraph_format.space_after = Pt(8)
        continue

    # 8. Rumus LaTeX Display ($$...$$)
    if line.strip().startswith("$$") and line.strip().endswith("$$"):
        eq_text = line.strip()[2:-2].strip()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(eq_text)
        run.font.name = "Cambria Math"
        run.font.size = Pt(10.5)
        run.font.italic = True
        i += 1
        continue

    # 9. List Numbering / Bullet
    list_match = re.match(r'^(\d+\.|\*|\-)\s+(.*)$', line.strip())
    if list_match and not in_daftar_pustaka:
        marker = list_match.group(1)
        item_text = list_match.group(2)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.25)
        
        # Add marker
        run_m = p.add_run(marker + " ")
        run_m.font.name = "Times New Roman"
        run_m.font.size = Pt(10.5)
        run_m.font.bold = True if marker[0].isdigit() else False
        
        format_inline_runs(p, item_text, default_size=Pt(10.5))
        i += 1
        continue

    # 10. Daftar Pustaka (APA Hanging Indent)
    if in_daftar_pustaka:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        format_inline_runs(p, line.strip(), default_size=Pt(10))
        i += 1
        continue

    # 11. Paragraf Reguler
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15

    # Cek apakah paragraf abstrak
    if "Implementasi pemeringkatan desil" in line or "The implementation of welfare" in line or line.startswith("**Kata Kunci:") or line.startswith("**Keywords:"):
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.right_indent = Inches(0.3)
        format_inline_runs(p, line.strip(), default_size=Pt(10), default_italic=(not line.startswith("**Kata Kunci:")))
    else:
        p.paragraph_format.first_line_indent = Inches(0.35)
        format_inline_runs(p, line.strip(), default_size=Pt(11))

    i += 1

doc.save(DOCX_FILE)
print(f"[+] Berhasil mengonversi dan menyimpan naskah Artikel Ilmiah ke: {DOCX_FILE}")
