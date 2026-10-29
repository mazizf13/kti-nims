"""
Script: convert_skripsi_to_word.py
Deskripsi: Mengonversi naskah Skripsi IndoBERT vs BiLSTM (Markdown) ke Microsoft Word (.docx)
           dengan standar tata tulis ilmiah Indonesia (Times New Roman, spasi 1.15, tabel terformat,
           dan daftar pustaka APA hanging indent).
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
SKRIPSI_DIR = os.path.join(BASE_DIR, "00_Skripsi_IndoBERT_vs_BiLSTM", "06_Naskah_Skripsi")
MD_FILE = os.path.join(SKRIPSI_DIR, "Skripsi_IndoBERT_vs_BiLSTM.md")
DOCX_FILE = os.path.join(SKRIPSI_DIR, "Skripsi_IndoBERT_vs_BiLSTM.docx")

print("[*] Membaca naskah Markdown Skripsi...")
with open(MD_FILE, "r", encoding="utf-8") as f:
    md_content = f.read()

doc = docx.Document()

# Page Setup: Margin 1 inch (2.54 cm)
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_table(table, num_cols):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    header_tr = table.rows[0]._tr.get_or_add_trPr()
    header_tr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    
    col_widths = None
    if num_cols == 5:
        # Tabel 3
        col_widths = [Inches(1.2), Inches(2.6), Inches(1.3), Inches(1.4), Inches(1.3)]
    elif num_cols == 6:
        # Tabel 1 atau Tabel 2
        col_widths = [Inches(1.8), Inches(1.0), Inches(1.1), Inches(1.0), Inches(1.4), Inches(1.2)]

    # Header Row
    for c_idx, cell in enumerate(table.rows[0].cells):
        tcPr = cell._tc.get_or_add_tcPr()
        tcPr.append(parse_xml(r'<w:shd %s w:fill="1F4E79"/>' % nsdecls('w')))
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        if col_widths and c_idx < len(col_widths):
            cell.width = col_widths[c_idx]
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(9.5)
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row in enumerate(table.rows[1:], start=1):
        shd_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcPr.append(parse_xml(r'<w:shd %s w:fill="%s"/>' % (nsdecls('w'), shd_color)))
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
            for p in cell.paragraphs:
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                elif c_idx == 1 and num_cols == 5:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.name = "Times New Roman"
                    run.font.size = Pt(9.0)
                    run.font.color.rgb = RGBColor(30, 41, 59)

lines = md_content.splitlines()
i = 0
in_daftar_pustaka = False

while i < len(lines):
    line = lines[i].rstrip()
    
    # 1. Tabel Markdown
    if line.strip().startswith("|") and line.strip().endswith("|"):
        current_table_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip().startswith("|"):
            current_table_lines.append(lines[i].strip())
            i += 1
            
        if len(current_table_lines) >= 3 and "---" in current_table_lines[1]:
            headers = [c.strip() for c in current_table_lines[0].split("|")[1:-1]]
            rows = []
            for r_line in current_table_lines[2:]:
                cols = [c.strip() for c in r_line.split("|")[1:-1]]
                cols_clean = [
                    re.sub(r"\*\*([^*]+)\*\*", r"\1", c)
                    .replace(r"\rightarrow", "->")
                    .replace(r"\\rightarrow", "->")
                    .replace(r"\$", "$")
                    for c in cols
                ]
                rows.append(cols_clean)
                
            tbl = doc.add_table(rows=len(rows)+1, cols=len(headers))
            for c_idx, h in enumerate(headers):
                clean_h = re.sub(r"\*\*([^*]+)\*\*", r"\1", h).replace(r"\rightarrow", "->").replace(r"\$", "$")
                tbl.cell(0, c_idx).text = clean_h
            for r_idx, row_vals in enumerate(rows, start=1):
                for c_idx, val in enumerate(row_vals):
                    if c_idx < len(headers):
                        tbl.cell(r_idx, c_idx).text = val
            style_table(tbl, len(headers))
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_after = Pt(6)
        continue

    # 2. Judul Dokumen (# Judul)
    if line.startswith("# "):
        title_text = line[2:].strip()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(title_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(15, 37, 55)

    # 3. Heading 2 (## BAB 1. PENDAHULUAN)
    elif line.startswith("## "):
        h2_text = line[3:].strip()
        if "DAFTAR PUSTAKA" in h2_text.upper():
            in_daftar_pustaka = True
        else:
            in_daftar_pustaka = False
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(h2_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(31, 78, 121)

    # 4. Heading 3 (### 1.1 Latar Belakang)
    elif line.startswith("### "):
        h3_text = line[4:].strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(h3_text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(40, 40, 40)

    # 5. Caption Tabel (*Tabel X...*)
    elif (line.startswith("*Tabel ") or line.startswith("*Gambar ")) and line.endswith("*"):
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if line.startswith("*Tabel "):
            p_cap.paragraph_format.space_before = Pt(8)
            p_cap.paragraph_format.space_after = Pt(3)
        else:
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(10)
        run = p_cap.add_run(line.strip("*"))
        run.font.name = "Times New Roman"
        run.font.size = Pt(9.5)
        run.font.italic = True
        run.font.bold = True
        run.font.color.rgb = RGBColor(70, 70, 70)

    # 6. Metadata Penulis
    elif line.startswith("**Penulis:**") or line.startswith("**Institusi") or line.startswith("**Email:"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        clean_meta = re.sub(r"\*\*([^*]+)\*\*", r"\1", line)
        run = p.add_run(clean_meta)
        run.font.name = "Times New Roman"
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(60, 60, 60)

    # 7. Pembatas
    elif line.strip() == "---":
        pass

    # 8. Paragraf atau List
    elif line.strip():
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        
        if in_daftar_pustaka:
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            text_run = line.strip()
        elif line.strip().startswith("* ") or line.strip().startswith("- "):
            p.paragraph_format.left_indent = Inches(0.25)
            text_run = line.strip()[2:]
        elif re.match(r"^\d+\.\s+", line.strip()):
            p.paragraph_format.left_indent = Inches(0.25)
            text_run = line.strip()
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            text_run = line.strip()
            
        text_run = re.sub(r"\$([^$]+)\$", r"\1", text_run)
        parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text_run)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.font.bold = True
            elif part.startswith("*") and part.endswith("*"):
                r = p.add_run(part[1:-1])
                r.font.italic = True
            elif part.startswith("`") and part.endswith("`"):
                r = p.add_run(part[1:-1])
                r.font.name = "Consolas"
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(180, 40, 40)
            else:
                clean_text = (
                    part.replace(r"\rightarrow", "->")
                    .replace(r"\\rightarrow", "->")
                    .replace(r"\sum", "sum")
                    .replace(r"\text", "")
                    .replace(r"\{", "{")
                    .replace(r"\}", "}")
                    .replace("\\", "")
                )
                r = p.add_run(clean_text)
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)

    i += 1

doc.save(DOCX_FILE)
print(f"[+] Berhasil mengonversi dan menyimpan naskah Skripsi ke: {DOCX_FILE}")
